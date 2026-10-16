from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from a PDF or DOCX file held in memory."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                return "\n".join(page.get_text("text") for page in document).strip()
        except Exception as exc:
            raise ValueError("Could not read this PDF. Make sure it is a valid PDF file.") from exc

    if suffix == ".docx":
        try:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            # Include table content because many resumes use tables for layout.
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(paragraphs).strip()
        except Exception as exc:
            raise ValueError("Could not read this DOCX file. Make sure it is a valid Word document.") from exc

    raise ValueError("Unsupported file type. Upload a PDF or DOCX file.")
