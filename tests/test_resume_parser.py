from docx import Document
from io import BytesIO

from backend.utils.resume_parser import extract_resume_text

def test_extract_docx_text_and_table():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Alex Example")
    document.add_paragraph("Skills: Python, SQL")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "Career matcher"
    document.save(buffer)

    result = extract_resume_text("resume.docx", buffer.getvalue())
    assert "Alex Example" in result
    assert "Python, SQL" in result
    assert "Career matcher" in result


def test_unsupported_file_type():
    try:
        extract_resume_text("resume.txt", b"some text")
    except ValueError as exc:
        assert "Unsupported file type" in str(exc)
    else:
        raise AssertionError("Expected ValueError")
